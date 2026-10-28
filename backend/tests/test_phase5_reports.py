"""
tests/test_phase5_reports.py

Tests for Excel and PDF report exports, formula injection sanitization,
HTTP content headers, and authorization.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.db.seed_phase4 import seed_phase4
from app.models.academic import AcademicYear
from app.services.report_service import ReportService, sanitize_excel_cell



class TestFormulaInjectionProtection:
    def test_sanitize_formula_injection_characters(self):
        """Values starting with dangerous formula injection characters are escaped."""
        assert sanitize_excel_cell("=1+1") == "'=1+1"
        assert sanitize_excel_cell("+SUM(A1:A10)") == "'+SUM(A1:A10)"
        assert sanitize_excel_cell("-10") == "'-10"
        assert sanitize_excel_cell("@SUM(A1:A5)") == "'@SUM(A1:A5)"
        assert sanitize_excel_cell("\tTABBED") == "'\tTABBED"
        assert sanitize_excel_cell("\rRETURN") == "'\rRETURN"

    def test_benign_values_unchanged(self):
        """Standard strings and numbers are unchanged."""
        assert sanitize_excel_cell("Computer Science") == "Computer Science"
        assert sanitize_excel_cell(42) == 42
        assert sanitize_excel_cell(99.5) == 99.5
        assert sanitize_excel_cell(None) is None


class TestReportGeneration:
    def test_excel_generation_service(self, db_session: Session):
        """Excel generation produces valid readable xlsx workbook."""
        seed_phase4(db_session)
        ay = AcademicYear(
            name="2026-2027",
            start_date=datetime(2026, 7, 1, tzinfo=timezone.utc).date(),
            end_date=datetime(2027, 6, 30, tzinfo=timezone.utc).date(),
            is_active=True,
        )

        db_session.add(ay)
        db_session.commit()

        service = ReportService(db_session)
        excel_buf = service.generate_excel_readiness_report(
            academic_year_id=ay.id,
            department="Computer Science and Engineering",
        )

        assert isinstance(excel_buf, io.BytesIO)
        excel_buf.seek(0)
        wb = openpyxl.load_workbook(excel_buf)
        assert "Readiness Summary" in wb.sheetnames
        assert "Criteria & Requirements" in wb.sheetnames
        assert "Gaps & Corrective Actions" in wb.sheetnames
        assert "Evidence Inventory" in wb.sheetnames

    def test_pdf_generation_service(self, db_session: Session):
        """PDF generation produces valid non-empty PDF document."""
        seed_phase4(db_session)
        ay = AcademicYear(
            name="2026-2027",
            start_date=datetime(2026, 7, 1, tzinfo=timezone.utc).date(),
            end_date=datetime(2027, 6, 30, tzinfo=timezone.utc).date(),
            is_active=True,
        )

        db_session.add(ay)
        db_session.commit()

        service = ReportService(db_session)
        pdf_buf = service.generate_pdf_readiness_report(
            academic_year_id=ay.id,
            department="Computer Science and Engineering",
        )

        assert isinstance(pdf_buf, io.BytesIO)
        content = pdf_buf.getvalue()
        assert len(content) > 100
        assert content.startswith(b"%PDF")


class TestReportAPIEndpoints:
    def test_excel_download_api(
        self,
        db_client: TestClient,
        db_session: Session,
        hod_headers: dict,
    ):
        """GET /api/v1/reports/excel streams xlsx with correct headers."""
        seed_phase4(db_session)
        ay = AcademicYear(
            name="2026-2027",
            start_date=datetime(2026, 7, 1, tzinfo=timezone.utc).date(),
            end_date=datetime(2027, 6, 30, tzinfo=timezone.utc).date(),
            is_active=True,
        )

        db_session.add(ay)
        db_session.commit()

        resp = db_client.get(
            f"/api/v1/reports/excel?academic_year_id={ay.id}&department=Computer Science and Engineering",
            headers=hod_headers,
        )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
        assert 'attachment; filename="NBA_Readiness_' in resp.headers["content-disposition"]
        assert len(resp.content) > 500

    def test_pdf_download_api(
        self,
        db_client: TestClient,
        db_session: Session,
        hod_headers: dict,
    ):
        """GET /api/v1/reports/pdf streams pdf with correct headers."""
        seed_phase4(db_session)
        ay = AcademicYear(
            name="2026-2027",
            start_date=datetime(2026, 7, 1, tzinfo=timezone.utc).date(),
            end_date=datetime(2027, 6, 30, tzinfo=timezone.utc).date(),
            is_active=True,
        )

        db_session.add(ay)
        db_session.commit()

        resp = db_client.get(
            f"/api/v1/reports/pdf?academic_year_id={ay.id}&department=Computer Science and Engineering",
            headers=hod_headers,
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert 'attachment; filename="NBA_Readiness_Summary_' in resp.headers["content-disposition"]
        assert resp.content.startswith(b"%PDF")

    def test_unauthenticated_request_rejected(self, db_client: TestClient):
        """Unauthenticated request to download reports is rejected with 401."""
        resp = db_client.get("/api/v1/reports/pdf")
        assert resp.status_code == 401
