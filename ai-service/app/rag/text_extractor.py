# text_extractor.py
"""
Document Text Extractor for NBA Evidence Files.

Supports multi-format ingestion:
- PDF files (.pdf) using pypdf
- Word documents (.docx) using python-docx
- Plain text and Markdown (.txt, .md)
- Data files (.csv, .json)
"""

import json
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional


def clean_extracted_text(raw_text: str) -> str:
    """
    Clean and normalize extracted text:
    - Replace multiple spaces and tabs with a single space
    - Normalize newline sequences
    - Remove non-printable control characters
    """
    if not raw_text:
        return ""

    # Remove non-printable control characters (keep standard newlines and tabs)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", raw_text)

    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Strip each line to remove leading/trailing horizontal spaces per line
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)

    # Replace 3+ consecutive newlines with 2 newlines
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def extract_from_pdf(file_path: str) -> Dict[str, Any]:
    """
    Extract text page-by-page from a PDF document using pypdf.
    """
    try:
        import pypdf

        reader = pypdf.PdfReader(file_path)
        pages = []
        full_text_parts = []

        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            cleaned = clean_extracted_text(page_text)
            pages.append({
                "page_number": idx + 1,
                "text": cleaned
            })
            if cleaned:
                full_text_parts.append(cleaned)

        full_text = "\n\n".join(full_text_parts)

        return {
            "status": "success",
            "file_type": "pdf",
            "total_pages": len(reader.pages),
            "full_text": full_text,
            "pages": pages
        }
    except Exception as error:
        return {
            "status": "error",
            "file_type": "pdf",
            "error": f"PDF extraction failed: {str(error)}",
            "full_text": "",
            "pages": []
        }


def extract_from_docx(file_path: str) -> Dict[str, Any]:
    """
    Extract text paragraphs and tables from a Word (.docx) document using python-docx.
    """
    try:
        import docx

        doc = docx.Document(file_path)
        paragraphs = []

        for p in doc.paragraphs:
            text = clean_extracted_text(p.text)
            if text:
                paragraphs.append(text)

        # Also extract table cells
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [clean_extracted_text(cell.text) for cell in row.cells]
                row_str = " | ".join(filter(None, row_cells))
                if row_str:
                    table_texts.append(row_str)

        all_text = "\n".join(paragraphs)
        if table_texts:
            all_text += "\n\n[Extracted Tables]\n" + "\n".join(table_texts)

        return {
            "status": "success",
            "file_type": "docx",
            "total_paragraphs": len(paragraphs),
            "full_text": all_text,
            "pages": [{"page_number": 1, "text": all_text}]
        }
    except Exception as error:
        return {
            "status": "error",
            "file_type": "docx",
            "error": f"DOCX extraction failed: {str(error)}",
            "full_text": "",
            "pages": []
        }


def extract_from_text(file_path: str) -> Dict[str, Any]:
    """
    Extract text from plain text or Markdown files (.txt, .md).
    """
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

        cleaned = clean_extracted_text(content)
        return {
            "status": "success",
            "file_type": "text",
            "full_text": cleaned,
            "pages": [{"page_number": 1, "text": cleaned}]
        }
    except Exception as error:
        return {
            "status": "error",
            "file_type": "text",
            "error": f"Text file extraction failed: {str(error)}",
            "full_text": "",
            "pages": []
        }


def extract_from_json(file_path: str) -> Dict[str, Any]:
    """
    Extract and serialize text from JSON files.
    """
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)

        text = json.dumps(data, indent=2, default=str)
        return {
            "status": "success",
            "file_type": "json",
            "full_text": text,
            "pages": [{"page_number": 1, "text": text}],
            "raw_data": data
        }
    except Exception as error:
        return {
            "status": "error",
            "file_type": "json",
            "error": f"JSON extraction failed: {str(error)}",
            "full_text": "",
            "pages": []
        }


def extract_document_text(file_path: str) -> Dict[str, Any]:
    """
    Automatically detect file format and extract text and metadata.
    """
    path = Path(file_path)
    if not path.exists():
        return {
            "status": "error",
            "error": f"File does not exist: {file_path}",
            "filename": path.name,
            "full_text": "",
            "pages": []
        }

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        result = extract_from_pdf(str(path))
    elif suffix in [".docx", ".doc"]:
        result = extract_from_docx(str(path))
    elif suffix in [".txt", ".md", ".markdown", ".csv"]:
        result = extract_from_text(str(path))
    elif suffix == ".json":
        result = extract_from_json(str(path))
    else:
        # Fallback to plain text reading
        result = extract_from_text(str(path))

    result["filename"] = path.name
    result["file_path"] = str(path.resolve())
    result["file_size_bytes"] = path.stat().st_size

    return result
